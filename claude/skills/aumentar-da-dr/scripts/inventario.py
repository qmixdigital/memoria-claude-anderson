# -*- coding: utf-8 -*-
"""Lista artigos candidatos a hospedar uma insercao de link, por host e tema.

    python inventario.py HOST "regex-do-tema" [--min-palavras 450] [--excluir-portais a,b,c] [--json saida.json]

HOST e um alias do ~/.ssh/config (opengravity, clinicas-vps, hostinger-vps-srv1166087).
O script se copia para o host, varre /srv/portais com find -L (boa parte dos
diretorios e symlink) e devolve, por artigo: portal, slug, categoria, data,
palavras e os dominios externos que o corpo ja linka.

Classificacao dos externos, para o vetting do passo 6 da skill:
  NEUTRO   americanas, fonte de noticia, orgao publico, periodico, rede social do portal
  CLIENTE  qualquer outro dominio, e TODO link de instagram.com (tier de cliente na rede)
Reprova tambem: dominio em --excluir-dominios ou em D:/PORTAIS/BACKLINKS/portais-bloqueados.txt,
portal com JSON gravado nos ultimos 90 min (outra sessao publicando), artigo que consta como
hospedeiro em qualquer planilha D:/PORTAIS/BACKLINKS/*.xlsx (guest post vendido), mais de 3
externos, texto curto. O que sobra ainda passa por impressao no GSC e leitura do tema.
"""
import argparse, io, json, os, re, shlex, subprocess, sys, tempfile

REMOTO = r'''
import json, re, os, subprocess, sys, time
RX = re.compile(sys.argv[1], re.I)
AGORA = time.time()
recentes = set()
MINW = int(sys.argv[2])
EXCL = set(x for x in sys.argv[3].split(",") if x)
files = subprocess.run("find -L /srv/portais -mindepth 3 -maxdepth 3 -path '*/data/*.json'", shell=True, capture_output=True, text=True).stdout.split()
out = []
for f in files:
    portal = f.split("/")[3]
    if portal.startswith("_") or portal in ("teste", "backups") or portal in EXCL: continue
    slug = os.path.basename(f)[:-5]
    try:
        if AGORA - os.path.getmtime(f) < 5400: recentes.add(portal)
        d = json.load(open(f, encoding="utf-8"))
    except Exception: continue
    if (d.get("status") or "publish") != "publish": continue
    c = d.get("content") or ""
    if not (RX.search(slug) or RX.search(d.get("title") or "")): continue
    ext = sorted(set(re.findall(r'href="https?://([^/"]+)', c)))
    out.append({"portal": portal, "slug": slug, "cat": (d.get("category") or {}).get("slug", ""),
                "date": (d.get("date") or "")[:10], "words": len(re.sub(r"<[^>]+>", " ", c).split()),
                "ext": ext, "title": (d.get("title") or "")[:100]})
sites = {x["slug"]: (x.get("baseUrl") or "https://" + x["domain"]).rstrip("/") for x in json.load(open("/opt/portal-engine/sites.json"))["sites"]}
print(json.dumps({"artigos": out, "recentes": sorted(recentes), "sites": sites}, ensure_ascii=False))
'''

NEUTRO = re.compile(r"americanas|amazon|mercadolivre|magazineluiza|shopee|acritica|globo|uol\.com|g1\.|cnn|bbc|estadao|folha\.|terra\.com|r7\.com|metropoles|correio|\.gov\.|\.org$|\.org\.br$|wikipedia|scielo|pubmed|plos|who\.int|confef\.org|chromewebstore|youtube|facebook|twitter|t\.co$|spotify|apple\.com|google\.com|microsoft")

def classifica(dominios, proprio=""):
    """Link absoluto para o proprio dominio do portal e interno, nao externo."""
    dominios = [d for d in dominios if d.replace("www.", "") != proprio]
    if any("instagram" in d for d in dominios): return "CLIENTE"
    return "CLIENTE" if any(not NEUTRO.search(d) for d in dominios) else "NEUTRO"

BLOCKLIST = "D:/PORTAIS/BACKLINKS/portais-bloqueados.txt"
BACKLINKS = "D:/PORTAIS/BACKLINKS"

def urls_vendidas():
    """URLs hospedeiras de TODAS as planilhas de cliente: artigo que esta nelas e guest post pago."""
    import glob
    try:
        import openpyxl
    except ImportError:
        return set()
    out = set()
    for x in glob.glob(os.path.join(BACKLINKS, "*.xlsx")):
        try:
            wb = openpyxl.load_workbook(x, read_only=True)
            ws = wb[wb.sheetnames[0]]
            for r in ws.iter_rows(min_row=2, values_only=True):
                if len(r) > 2 and isinstance(r[2], str) and r[2].startswith("http"):
                    out.add(r[2].rstrip("/").split("://", 1)[1].replace("www.", ""))
        except Exception:
            pass
    return out

def bloqueados():
    if not os.path.exists(BLOCKLIST): return set()
    return set(l.split("#")[0].strip() for l in io.open(BLOCKLIST, encoding="utf-8") if l.split("#")[0].strip())

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("host"); ap.add_argument("tema", help="regex aplicada ao slug e ao titulo")
    ap.add_argument("--min-palavras", type=int, default=450)
    ap.add_argument("--excluir-portais", default="", help="slugs de portal a pular")
    ap.add_argument("--excluir-dominios", default="", help="dominios a pular, separados por virgula (os que ja apontam para o site atendido, do ledger)")
    ap.add_argument("--limite", type=int, default=6, help="artigos por portal na tela (o --json traz todos)")
    ap.add_argument("--json", help="salvar a lista completa em JSON")
    ap.add_argument("--todos", action="store_true", help="mostrar tambem os reprovados")
    ap.add_argument("--recentes-ok", action="store_true", help="nao reprovar portal com gravacao recente; SO quando a gravacao foi desta mesma sessao (lote proprio ja terminado)")
    a = ap.parse_args()
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    with tempfile.NamedTemporaryFile("w", suffix=".py", delete=False, encoding="utf-8") as t:
        t.write(REMOTO); tmp = t.name
    subprocess.run(["scp", "-q", tmp, "%s:/tmp/inv_da_dr.py" % a.host], check=True)
    r = subprocess.run(["ssh", "-n", a.host, "python3 /tmp/inv_da_dr.py %s %d %s" % (shlex.quote(a.tema), a.min_palavras, shlex.quote(a.excluir_portais or ""))],
                       capture_output=True, text=True, encoding="utf-8", errors="replace")
    os.unlink(tmp)
    bruto = json.loads(r.stdout or '{"artigos":[],"recentes":[],"sites":{}}')
    dados, recentes, sites = bruto["artigos"], set(bruto["recentes"]), bruto["sites"]
    dom_de = {k: re.sub(r"^https?://(www\.)?", "", v) for k, v in sites.items()}
    excl_dom = set(d.strip().replace("www.", "") for d in a.excluir_dominios.split(",") if d.strip()) | bloqueados()
    vendidas = urls_vendidas()
    for x in dados:
        dom = dom_de.get(x["portal"], "")
        x["dominio"] = dom
        x["ext"] = [e for e in x["ext"] if e.replace("www.", "") != dom]
        x["classe"] = classifica(x["ext"], dom)
        x["motivo"] = ""
        if dom in excl_dom or x["portal"] in excl_dom: x["motivo"] = "dominio ja usado ou bloqueado"
        elif x["portal"] in recentes and not a.recentes_ok: x["motivo"] = "portal com JSON gravado nos ultimos 90 min (lote em andamento)"
        elif any(u.startswith(dom) and u.endswith("/" + x["slug"]) for u in vendidas): x["motivo"] = "guest post vendido a cliente (esta numa planilha .xlsx)"
        elif x["classe"] != "NEUTRO": x["motivo"] = "linka cliente"
        elif len(x["ext"]) > 3: x["motivo"] = "mais de 3 externos"
        elif x["words"] < a.min_palavras: x["motivo"] = "curto"
        x["aprovado"] = not x["motivo"]
    if a.json:
        json.dump(dados, open(a.json, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    por_portal = {}
    for x in dados:
        if x["aprovado"] or a.todos: por_portal.setdefault(x["portal"], []).append(x)
    print("%d artigos no tema, %d aprovados no vetting automatico. Falta conferir a mao: impressao da pagina no GSC (gsc_paginas.py --url) e tema plausivel." % (len(dados), sum(1 for x in dados if x["aprovado"])))
    if recentes: print("portais com gravacao recente (fora desta rodada):", ", ".join(sorted(recentes)))
    for p, xs in sorted(por_portal.items()):
        print("-- %s  %s  (%d)" % (p, dom_de.get(p, ""), len(xs)))
        for x in sorted(xs, key=lambda x: -x["words"])[:a.limite]:
            print("   %s %4dw %s %-60s ext=%s %s" % ("OK " if x["aprovado"] else "-- ", x["words"], x["date"], x["slug"][:60], ",".join(x["ext"])[:40], x["motivo"]))

if __name__ == "__main__":
    main()
